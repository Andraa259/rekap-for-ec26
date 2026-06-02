import streamlit as st
import pandas as pd
import io
import re
import plotly.express as px
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Wonderful Class Presence Hub", layout="wide")

st.title("📊 Wonderful Class Presence Dashboard & Dynamic Matrix")
st.write("Sistem Matriks Presensi Otomatis (No, Nama, NIM/NPM, Total Hadir, P1 - P-Dinamis) dengan Fitur Auto-Expand Kolom.")

# ----------------- SIDEBAR INPUT -----------------
st.sidebar.header("📁 Unggah Dokumen Presensi")

master_file = st.sidebar.file_uploader(
    "1. File Master / Rekap Sebelumnya (Opsional, kosongkan jika baru mulai pertama kali)", 
    type=["xlsx", "xls"],
    key="master_upload"
)

source_files = st.sidebar.file_uploader(
    "2. Unggah Form Absensi Harian Baru (Pertemuan Bebas/Dinamis)", 
    type=["xlsx", "xls"], 
    accept_multiple_files=True,
    key="sources_upload"
)

# Fungsi bantuan untuk mendeteksi kolom NIM/NPM/NBI secara fleksibel
def temukan_kolom_nim(df_columns):
    for col in df_columns:
        col_clean = str(col).strip().lower()
        if col_clean in ['nim', 'npm', 'nbi', 'nim / npm', 'nim/npm', 'no. induk']:
            return col
    return None

# ----------------- LOGIKA UTAMA -----------------
if st.sidebar.button("🚀 Proses & Sinkronisasi Presensi", type="primary"):
    if not source_files and master_file is None:
        st.error("❌ Mohon unggah minimal File Master Rekap Sebelumnya ATAU File Form Absensi Harian!")
    else:
        with st.spinner("Sedang memproses dan menyelaraskan data presensi secara dinamis..."):
            try:
                # 1. DETEKSI PERTEMUAN MAKSIMAL DARI FILE BARU YANG DIUNGGAH
                max_pertemuan_baru = 0
                file_info_list = []
                
                for src in source_files:
                    nama_file = src.name.lower()
                    angka_terdeteksi = re.findall(r'\d+', nama_file)
                    if angka_terdeteksi:
                        nomor_p = int(angka_terdeteksi[0])
                        max_pertemuan_baru = max(max_pertemuan_baru, nomor_p)
                        file_info_list.append((src, nomor_p))
                    else:
                        st.warning(f"⚠️ Gagal mendeteksi nomor pertemuan pada file '{src.name}'. Pastikan ada angka pertemuan di nama file.")

                # 2. BACA MASTER DATA / REKAP SEBELUMNYA
                df_rekap = pd.DataFrame()
                existing_meetings = []
                max_pertemuan_lama = 0
                
                if master_file is not None:
                    df_master_raw = pd.read_excel(master_file)
                    df_master_raw.columns = df_master_raw.columns.str.strip()
                    
                    if df_master_raw.columns[0].startswith('Unnamed:'):
                        df_master_raw = df_master_raw.iloc[:, 1:]
                    
                    kolom_nim_master = temukan_kolom_nim(df_master_raw.columns)
                    if not kolom_nim_master or 'Nama' not in df_master_raw.columns:
                        st.error("❌ File Master / Rekap Sebelumnya harus memiliki kolom 'Nama' dan kolom Identitas (NIM / NPM / NBI)!")
                        st.stop()
                        
                    df_master_raw[kolom_nim_master] = df_master_raw[kolom_nim_master].astype(str).str.strip()
                    df_master_raw['Nama'] = df_master_raw['Nama'].astype(str).str.strip()
                    
                    df_rekap['Nama'] = df_master_raw['Nama'].values
                    df_rekap['NIM / NPM'] = df_master_raw[kolom_nim_master].values
                    
                    # Deteksi kolom pertemuan yang sudah ada di file rekap lama
                    for col in df_master_raw.columns:
                        if col.startswith("Pertemuan "):
                            existing_meetings.append(col)
                            angka_p = int(col.replace("Pertemuan ", ""))
                            max_pertemuan_lama = max(max_pertemuan_lama, angka_p)
                            df_rekap[col] = df_master_raw[col].fillna("Alpa").values
                            
                    st.info(f"📂 Melanjutkan rekap sebelumnya. Terdeteksi {len(df_rekap)} anggota dan {len(existing_meetings)} pertemuan lama.")
                else:
                    # JIKA BARU PERTAMA KALI (BUILT-IN DATABASE OTOMATIS)
                    list_for_built_in = []
                    for src in source_files:
                        try:
                            df_check = pd.read_excel(src)
                            df_check.columns = df_check.columns.str.strip()
                            kolom_nim_src = temukan_kolom_nim(df_check.columns)
                            if kolom_nim_src and 'Nama' in df_check.columns:
                                df_temp = df_check[['Nama', kolom_nim_src]].copy()
                                df_temp.columns = ['Nama', 'NIM / NPM']
                                list_for_built_in.append(df_temp)
                        except:
                            continue
                    
                    if list_for_built_in:
                        df_built_in = pd.concat(list_for_built_in, ignore_index=True)
                        df_built_in['NIM / NPM'] = df_built_in['NIM / NPM'].astype(str).str.strip()
                        df_built_in['Nama'] = df_built_in['Nama'].astype(str).str.strip()
                        
                        df_base_members = df_built_in.drop_duplicates(subset=['NIM / NPM']).sort_values(by=['Nama']).reset_index(drop=True)
                        df_rekap['Nama'] = df_base_members['Nama'].values
                        df_rekap['NIM / NPM'] = df_base_members['NIM / NPM'].values
                        st.success(f"✨ File Master kosong. Sistem otomatis membuat format baru dari form harian: Terdeteksi {len(df_rekap)} anggota unik.")
                    else:
                        st.error("❌ Gagal memproses data. Pastikan file form harian memiliki kolom 'Nama' dan kolom Identitas seperti 'NIM / NPM' atau 'NBI'!")
                        st.stop()

                # 3. TENTUKAN JUMLAH KOLOM PERTEMUAN SECARA DINAMIS
                total_maksimal_pertemuan = max(max_pertemuan_lama, max_pertemuan_baru)
                
                if total_maksimal_pertemuan == 0:
                    st.error("❌ Tidak ada kolom pertemuan valid yang terdeteksi dari nama file.")
                    st.stop()
                
                list_pertemuan_all = [f"Pertemuan {i}" for i in range(1, total_maksimal_pertemuan + 1)]
                
                for p in list_pertemuan_all:
                    if p not in df_rekap.columns:
                        df_rekap[p] = "Alpa"

                # 4. ISI/UPDATE DATA DARI FILE FORM ABSENSI BARU
                kehadiran_per_hari = {}
                
                for src, nomor_p in file_info_list:
                    df_src = pd.read_excel(src)
                    df_src.columns = df_src.columns.str.strip()
                    
                    kolom_nim_src = temukan_kolom_nim(df_src.columns)
                    if not kolom_nim_src:
                        st.warning(f"⚠️ File '{src.name}' dilewati karena tidak ditemukan kolom identitas NIM/NPM/NBI.")
                        continue
                        
                    df_src[kolom_nim_src] = df_src[kolom_nim_src].astype(str).str.strip()
                    df_src = df_src.drop_duplicates(subset=[kolom_nim_src], keep='first')
                    
                    target_kolom = f"Pertemuan {nomor_p}"
                    
                    # Update status Hadir
                    status_hadir_baru = df_rekap['NIM / NPM'].isin(df_src[kolom_nim_src])
                    df_rekap.loc[status_hadir_baru, target_kolom] = 'Hadir'

                # 5. HITUNG TOTAL HADIR AKTUAL AKHIR
                df_rekap['Total Hadir'] = (df_rekap[list_pertemuan_all] == 'Hadir').sum(axis=1)
                
                # Susun struktur susunan kolom final: No. -> Nama -> NIM / NPM -> Total Hadir -> Sesi Pertemuan
                df_rekap.insert(0, 'No.', range(1, len(df_rekap) + 1))
                susunan_kolom_final = ['No.', 'Nama', 'NIM / NPM', 'Total Hadir'] + list_pertemuan_all
                df_rekap = df_rekap[susunan_kolom_final]

                for p in list_pertemuan_all:
                    kehadiran_per_hari[p] = (df_rekap[p] == 'Hadir').sum()

                # --- 🔍 FITUR TAMBAHAN: PENGECEKAN KESAMAAN NAMA (DUPLIKAT) 🔍 ---
                # Mengabaikan huruf besar/kecil (case-insensitive) untuk akurasi pengecekan nama kembar
                df_nama_lower = df_rekap['Nama'].str.lower()
                duplikat_nama = df_rekap[df_nama_lower.duplicated(keep=False)]
                
                if not duplikat_nama.empty:
                    st.sidebar.markdown("---")
                    st.sidebar.warning("⚠️ **Perhatian: Terdeteksi Kesamaan Nama!**")
                    # Kelompokkan nama yang sama untuk ditampilkan di Sidebar Note
                    for nama_tertentu, group in duplikat_nama.groupby(df_rekap['Nama'].str.lower()):
                        nama_asli = group['Nama'].iloc[0]
                        nim_list = ", ".join(group['NIM / NPM'].tolist())
                        st.sidebar.write(f"- Nama **\"{nama_asli}\"** ditemukan lebih dari satu kali dengan NIM/NPM berbeda: ({nim_list})")
                    st.sidebar.info("💡 *Sistem tetap berjalan normal karena pencocokan mutlak menggunakan NIM / NPM.*")

                # ----------------- VISUALISASI DASHBOARD -----------------
                st.subheader("📌 Ringkasan Metrik Kelas (Dinamis)")
                m1, m2, m3 = st.columns(3)
                m1.metric("Total Anggota Terdata", f"{len(df_rekap)} Orang")
                m2.metric("Total Sesi Saat Ini", f"{total_maksimal_pertemuan} Pertemuan")
                
                total_maksimal_hadir = len(df_rekap) * total_maksimal_pertemuan
                persentase_total = ((df_rekap['Total Hadir'].sum() / total_maksimal_hadir) * 100).round(1)
                m3.metric("Rata-rata Presensi Kelas Overall", f"{persentase_total} %")
                
                st.markdown("---")
                st.subheader(f"📈 Visualisasi Grafik Batang Kehadiran (P1 - P{total_maksimal_pertemuan})")
                df_grafik = pd.DataFrame(list(kehadiran_per_hari.items()), columns=['Pertemuan', 'Jumlah Hadir'])
                fig = px.bar(
                    df_grafik, x='Pertemuan', y='Jumlah Hadir', text='Jumlah Hadir',
                    color_discrete_sequence=['#1E3A8A']
                )
                fig.update_traces(textposition='outside')
                st.plotly_chart(fig, use_container_width=True)

                st.markdown("---")
                st.subheader("👀 Preview Tabel Matriks Terformat")
                st.dataframe(df_rekap, use_container_width=True)

                # --- EXCEL FORMATTING WITH OPENPYXL (STRUKTUR KETAT PERMINTAAN) ---
                wb = Workbook()
                ws = wb.active
                ws.title = "Rekap Presensi Wonderful"
                ws.views.sheetView[0].showGridLines = True

                columns_list = df_rekap.columns.tolist()
                total_cols = len(columns_list)
                last_col_letter = get_column_letter(1 + total_cols)

                # 1. Judul Utama di Baris 2 & 3
                ws.merge_cells(f"B2:{last_col_letter}2")
                ws.merge_cells(f"B3:{last_col_letter}3")
                ws["B2"] = "REKAPITULASI PRESENSI"
                ws["B3"] = "WONDERFUL CLASS 2026"

                title_font = Font(name="Calibri", size=14, bold=True)
                center_alignment = Alignment(horizontal="center", vertical="center")
                
                ws["B2"].font = title_font
                ws["B2"].alignment = center_alignment
                ws["B3"].font = title_font
                ws["B3"].alignment = center_alignment
                ws.row_dimensions[2].height = 24
                ws.row_dimensions[3].height = 24

                # 2. Baris Judul Tabel / Header (Baris 6)
                header_row = 6
                ws.row_dimensions[header_row].height = 28
                
                thin_side = Side(border_style="thin", color="000000")
                thick_side = Side(border_style="medium", color="000000")
                
                data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
                header_border = Border(left=thin_side, right=thin_side, top=thick_side, bottom=thick_side)
                
                header_font = Font(name="Calibri", size=11, bold=True)
                header_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")

                for col_idx, col_name in enumerate(columns_list, start=2):
                    cell = ws.cell(row=header_row, column=col_idx)
                    cell.value = col_name
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_alignment
                    cell.border = header_border

                # 3. Pengisian Isi Data (Baris 7 Dst) -> Semua Sel Wajib Center & Border
                data_font = Font(name="Calibri", size=11)
                current_row = 7
                
                for _, row_data in df_rekap.iterrows():
                    ws.row_dimensions[current_row].height = 20
                    for col_idx, col_name in enumerate(columns_list, start=2):
                        cell = ws.cell(row=current_row, column=col_idx)
                        val = row_data[col_name]
                        cell.value = "" if pd.isna(val) else val
                        cell.font = data_font
                        cell.alignment = center_alignment
                        cell.border = data_border
                    current_row += 1

                # 4. Pengaturan Lebar Kolom (Kolom A Sengaja Dikosongkan)
                ws.column_dimensions['A'].width = 4
                for col_idx in range(2, 2 + total_cols):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(6, current_row))
                    ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

                # Ekspor file matriks ke memory untuk siap unduh
                output = io.BytesIO()
                wb.save(output)
                processed_data = output.getvalue()

                st.sidebar.markdown("---")
                st.sidebar.success("🎉 Matriks Presensi Berhasil Diperbarui!")
                st.sidebar.download_button(
                    label="📥 Download Excel Hasil Rekap Dinamis",
                    data=processed_data,
                    file_name="Rekap_Dinamis_WonderfulClass_2026.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            except Exception as e:
                st.error(f"Terjadi kesalahan teknis: {str(e)}")
